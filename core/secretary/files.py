from __future__ import annotations

import hashlib
import re
import shutil
import subprocess
import tempfile
import zipfile
from xml.etree import ElementTree
from datetime import date, datetime
from pathlib import Path

from .database import Database
from .memory import MemoryEngine
from .models import ModelRouter


class FilePipeline:
    IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp", ".heic", ".heif", ".bmp", ".gif"}
    VIDEO_SUFFIXES = {".mp4", ".mov", ".m4v", ".avi", ".mkv", ".webm"}
    DOCUMENT_SUFFIXES = {
        ".txt", ".md", ".csv", ".json", ".log", ".yaml", ".yml", ".xml", ".html", ".htm",
        ".ics", ".vcf", ".rtf", ".eml", ".pdf", ".doc", ".docx", ".xlsx", ".xlsm",
        ".pptx", ".ods", ".zip",
    }
    RETRY_MARKERS = ("请先在设置中配置", "模型分析失败", "关键帧分析失败", "解析失败", "解析组件未安装")

    def __init__(self, files_root: Path, database: Database, memory: MemoryEngine, models: ModelRouter):
        self.files_root, self.database, self.memory, self.models = files_root, database, memory, models

    def ingest(self, source: Path, original_name: str = "") -> dict:
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        existing = self.database.rows("SELECT * FROM files WHERE hash=?", (digest,))
        if existing:
            record = existing[0]
            can_reanalyze = Path(record["stored_path"]).suffix.lower() in self.IMAGE_SUFFIXES | self.VIDEO_SUFFIXES | self.DOCUMENT_SUFFIXES
            failed_before = not record["extracted_text"].strip() or any(marker in record["extracted_text"] for marker in self.RETRY_MARKERS)
            if can_reanalyze and failed_before:
                result = self.reanalyze(digest)
                return {**result, "duplicate": True, "reanalyzed": True, "can_reanalyze": True}
            return {**record, "duplicate": True, "reanalyzed": False, "can_reanalyze": can_reanalyze}
        month = date.today().strftime("%Y-%m")
        target_dir = self.files_root / month
        target_dir.mkdir(parents=True, exist_ok=True)
        original_name = original_name or source.name
        safe_name = original_name.replace("/", "_").replace("\\", "_")
        target = target_dir / f"{digest[:10]}-{safe_name}"
        shutil.copy2(source, target)
        text = self.extract(target)
        relative = str(target.relative_to(self.files_root.parent))
        with self.database.connect() as db:
            db.execute("INSERT INTO files(hash,original_name,stored_path,created_at,extracted_text) VALUES(?,?,?,?,?)", (digest, original_name, relative, datetime.now().isoformat(timespec="seconds"), text))
        self.memory.append_daily(f"收录文件 [{original_name}](../../{relative})：{text[:300] or '暂无可提取文字'}", "文件上传", inferred=not bool(text))
        return {"hash": digest, "original_name": original_name, "stored_path": relative, "extracted_text": text, "duplicate": False, "reanalyzed": False, "can_reanalyze": target.suffix.lower() in self.IMAGE_SUFFIXES | self.VIDEO_SUFFIXES | self.DOCUMENT_SUFFIXES}

    def reanalyze(self, digest: str) -> dict:
        rows = self.database.rows("SELECT * FROM files WHERE hash=?", (digest,))
        if not rows:
            raise KeyError("文件不存在")
        record = rows[0]
        path = (self.files_root.parent / record["stored_path"]).resolve()
        root = self.files_root.parent.resolve()
        if root not in path.parents or not path.is_file():
            raise FileNotFoundError("原始文件不存在")
        if path.suffix.lower() not in self.IMAGE_SUFFIXES | self.VIDEO_SUFFIXES | self.DOCUMENT_SUFFIXES:
            raise ValueError("这种文件格式暂不支持重新分析")
        text = self.extract(path)
        with self.database.connect() as db:
            db.execute("UPDATE files SET extracted_text=? WHERE hash=?", (text, digest))
        self.memory.append_daily(f"重新分析文件 [{record['original_name']}](../../{record['stored_path']})：{text[:300]}", "文件重新分析", inferred=False)
        return {**record, "extracted_text": text, "duplicate": True, "reanalyzed": True, "can_reanalyze": True}

    def delete(self, digest: str) -> dict:
        rows = self.database.rows("SELECT * FROM files WHERE hash=?", (digest,))
        if not rows:
            raise KeyError("文件不存在")
        record = rows[0]
        path = (self.files_root.parent / record["stored_path"]).resolve()
        files_root = self.files_root.resolve()
        if files_root not in path.parents:
            raise ValueError("拒绝删除归档目录之外的文件")
        path.unlink(missing_ok=True)
        with self.database.connect() as db:
            db.execute("DELETE FROM files WHERE hash=?", (digest,))
        return {"ok": True, "hash": digest, "original_name": record["original_name"]}

    def extract(self, path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md", ".csv", ".json", ".log", ".yaml", ".yml", ".ics", ".vcf"}:
            return self._read_text(path)[:200000]
        if suffix in {".html", ".htm"}:
            from .connectors import html_to_text
            return html_to_text(self._read_text(path))[:200000]
        if suffix == ".xml":
            try:
                return "\n".join(text.strip() for text in ElementTree.parse(path).getroot().itertext() if text.strip())[:200000]
            except ElementTree.ParseError:
                return self._read_text(path)[:200000]
        if suffix == ".rtf":
            text = self._read_text(path)
            text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
            text = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", text)
            return re.sub(r"[{}]", "", text)[:200000]
        if suffix == ".eml":
            from .connectors import parse_message
            message = parse_message(path.read_bytes())
            names = "、".join(item["name"] for item in message["attachments"])
            return (
                f"主题：{message['subject']}\n发件人：{message['from']}\n日期：{message['date']}\n\n{message['body']}"
                + (f"\n\n附件：{names}" if names else "")
            )[:200000]
        if suffix == ".pdf":
            try:
                from pypdf import PdfReader
                return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)[:200000]
            except ImportError:
                return "PDF 解析组件未安装，请安装 documents 可选依赖"
        if suffix == ".doc":
            return self._legacy_word_text(path)
        if suffix == ".docx":
            try:
                from docx import Document
                return "\n".join(p.text for p in Document(str(path)).paragraphs)[:200000]
            except ImportError:
                return "Word 解析组件未安装，请安装 documents 可选依赖"
        if suffix in {".xlsx", ".xlsm"}:
            try:
                from openpyxl import load_workbook
                book = load_workbook(path, read_only=True, data_only=True)
                return "\n".join(",".join("" if v is None else str(v) for v in row) for sheet in book for row in sheet.iter_rows(values_only=True))[:200000]
            except ImportError:
                return "Excel 解析组件未安装，请安装 documents 可选依赖"
        if suffix == ".pptx":
            return self._archive_xml_text(path, "ppt/slides/slide")
        if suffix == ".ods":
            return self._archive_xml_text(path, "content.xml")
        if suffix == ".zip":
            with zipfile.ZipFile(path) as archive:
                names = [name for name in archive.namelist() if not name.startswith("__MACOSX/")]
            return "压缩包内容：\n" + "\n".join(names[:500])
        if suffix in self.IMAGE_SUFFIXES:
            if not self.models.route("vision"):
                return "图片已归档；请先在设置中配置“图片处理”模型，之后可重新上传分析。"
            try:
                return self.models.analyze_images("vision", [path], "请用中文识别并总结这张图片中的有用信息。重点提取票据金额、日期、商户，体检指标，日程、地址和人物关系；明确区分事实与推测。")
            except Exception as exc:
                return f"图片已归档，但模型分析失败：{exc}"
        if suffix in self.VIDEO_SUFFIXES:
            task = "video" if self.models.route("video") else "vision"
            if not self.models.route(task):
                return "视频已归档；请先在设置中配置“视频处理”或“图片处理”模型，之后可重新上传分析。"
            try:
                frames = self._video_frames(path)
                try:
                    return self.models.analyze_images(task, frames, "这些图片按时间顺序来自同一段视频。请用中文总结视频内容、关键事件、可见文字、时间地点和可能值得记录的事项；不要把不确定内容说成事实。")
                finally:
                    shutil.rmtree(frames[0].parent, ignore_errors=True)
            except Exception as exc:
                return f"视频已归档，但关键帧分析失败：{exc}"
        return ""

    @staticmethod
    def _read_text(path: Path) -> str:
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "gb18030", "utf-16"):
            try:
                return raw.decode(encoding)
            except UnicodeDecodeError:
                continue
        return raw.decode("utf-8", errors="replace")

    @staticmethod
    def _archive_xml_text(path: Path, member_prefix: str) -> str:
        try:
            with zipfile.ZipFile(path) as archive:
                names = sorted(name for name in archive.namelist() if name.startswith(member_prefix) and name.endswith(".xml"))
                sections = []
                for name in names:
                    root = ElementTree.fromstring(archive.read(name))
                    text = " ".join(value.strip() for value in root.itertext() if value.strip())
                    if text:
                        sections.append(text)
            return "\n".join(sections)[:200000]
        except (zipfile.BadZipFile, ElementTree.ParseError, KeyError) as exc:
            return f"文档解析失败：{exc}"

    @staticmethod
    def _legacy_word_text(path: Path) -> str:
        commands = []
        if shutil.which("textutil"):
            commands.append(["textutil", "-convert", "txt", "-stdout", str(path)])
        if shutil.which("antiword"):
            commands.append(["antiword", str(path)])
        if not commands:
            return "旧版 Word 解析组件未安装；请将 .doc 另存为 .docx 后重新上传"
        errors = []
        for command in commands:
            try:
                result = subprocess.run(command, check=True, capture_output=True, timeout=60)
                text = result.stdout.decode("utf-8", errors="replace").strip()
                if text:
                    return text[:200000]
            except (subprocess.SubprocessError, OSError) as exc:
                errors.append(str(exc))
        detail = errors[-1] if errors else "没有提取到文字"
        return f"旧版 Word 解析失败：{detail}；可将 .doc 另存为 .docx 后重试"

    @staticmethod
    def _video_frames(path: Path) -> list[Path]:
        try:
            import imageio_ffmpeg
        except ImportError as exc:
            raise RuntimeError("缺少 imageio-ffmpeg，请重新安装 requirements.txt") from exc
        frame_dir = Path(tempfile.mkdtemp(prefix="second-brain-frames-"))
        pattern = frame_dir / "frame-%02d.png"
        command = [imageio_ffmpeg.get_ffmpeg_exe(), "-hide_banner", "-loglevel", "error", "-i", str(path), "-vf", "fps=1/10,scale=1280:-2", "-frames:v", "6", str(pattern)]
        subprocess.run(command, check=True, timeout=180)
        frames = sorted(frame_dir.glob("frame-*.png"))
        if not frames:
            fallback = frame_dir / "frame-01.png"
            subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(), "-hide_banner", "-loglevel", "error", "-i", str(path), "-frames:v", "1", "-vf", "scale=1280:-2", str(fallback)], check=True, timeout=180)
            frames = [fallback] if fallback.exists() else []
        if not frames:
            shutil.rmtree(frame_dir, ignore_errors=True)
            raise RuntimeError("无法从视频中提取画面")
        return frames
